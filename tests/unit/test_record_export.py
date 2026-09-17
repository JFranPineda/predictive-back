"""The workbook has to be the sheet the customer already reads."""

from io import BytesIO

import openpyxl
import pytest

from modules.measurements.infrastructure.record_export import (
    ExportBand,
    ExportBlock,
    ExportColumn,
    build_workbook,
)

EQUIPMENT = {
    "id": 1,
    "name": "MOTOR",
    "tag": "MB101001A",
    "group": "BBA. AGUA CRUDA - TAG:A",
    "area": "101 - ETA",
}

COLUMNS = [
    ExportColumn(date="2013-02-08", order_code="MPd-AV-N°001-13", operators="CT",
                 operating={"rpm": "1630", "freq_hz": "54"}),
    ExportColumn(date="2013-03-02", order_code="MPd-AV-N°002-13", operators="CT / WG",
                 operating={"rpm": "1613", "freq_hz": "54"}),
]

BLOCK = ExportBlock(
    title="Velocidad vibracional",
    unit="mm/s",
    aggregation="rms",
    decimals=2,
    rows=(
        {"component": "MOTOR", "side_label": "LADO LIBRE", "label": "1H",
         "values": ["0.70", "1.30"], "missing": ["", ""]},
        {"component": "MOTOR", "side_label": "LADO LIBRE", "label": "1V",
         "values": ["0.80", "0.70"], "missing": ["", ""]},
        {"component": "MOTOR", "side_label": "LADO ACOPLE", "label": "2H",
         "values": ["1.00", None], "missing": ["", "--"]},
        {"component": "BOMBA", "side_label": "LADO ACOPLE", "label": "3H",
         "values": ["2.20", "2.10"], "missing": ["", ""]},
    ),
)


@pytest.fixture(scope="module")
def sheet():
    payload = build_workbook(
        equipment=EQUIPMENT,
        columns=COLUMNS,
        blocks=[BLOCK],
        bands=[ExportBand(status="Alarma", minimum="4.5", maximum="7.1")],
        operating_labels={"rpm": "Velocidad (rpm)", "freq_hz": "Frecuencia (Hz)"},
        title="SERVICIO DE MANTENIMIENTO PREDICTIVO",
    )
    return openpyxl.load_workbook(BytesIO(payload)).active


def cells_of(sheet, label: str) -> list:
    for row in range(1, sheet.max_row + 1):
        if str(sheet.cell(row=row, column=1).value or "").startswith(label):
            return [sheet.cell(row=row, column=column).value for column in range(4, 6)]
    raise AssertionError(f"'{label}' not found")


class TestLayout:
    def test_the_sheet_is_named_after_the_equipment(self, sheet):
        assert sheet.title == "MB101001A"

    def test_one_column_per_round(self, sheet):
        assert cells_of(sheet, "FECHA") == ["2013-02-08", "2013-03-02"]

    def test_the_round_is_signed(self, sheet):
        # "LECTURAS TOMADAS POR: CT / WG" is how the source sheet signs it.
        assert cells_of(sheet, "LECTURAS TOMADAS POR") == ["CT", "CT / WG"]

    def test_the_limits_in_force_are_printed(self, sheet):
        # The reader has to see the criterion the values were graded against,
        # not just the values.
        band = next(
            row
            for row in range(1, sheet.max_row + 1)
            if str(sheet.cell(row=row, column=1).value or "").startswith("Alarma")
        )
        assert sheet.cell(row=band, column=2).value == "4.5"
        assert sheet.cell(row=band, column=3).value == "7.1"

    def test_running_conditions_close_the_sheet(self, sheet):
        assert cells_of(sheet, "Velocidad (rpm)") == [1630, 1613]

    def test_the_label_columns_stay_while_dates_scroll(self, sheet):
        assert sheet.freeze_panes == "D1"


class TestValues:
    def test_numbers_are_numbers_not_text(self, sheet):
        # A spreadsheet of strings cannot be charted or averaged, which is the
        # first thing the customer does with it.
        row = _row_of(sheet, "1H")
        assert sheet.cell(row=row, column=4).value == 0.7

    def test_a_round_that_could_not_be_measured_says_so(self, sheet):
        # Blank means the round does not exist for that point; "--" means it
        # exists and nothing could be read. The source sheet distinguishes
        # them and so does this.
        row = _row_of(sheet, "2H")
        assert sheet.cell(row=row, column=5).value == "--"

    def test_component_and_side_are_merged_like_the_printed_sheet(self, sheet):
        merges = {str(merge) for merge in sheet.merged_cells.ranges}
        first = _row_of(sheet, "1H")
        # MOTOR spans its three points; its free end spans two of them.
        assert f"A{first}:A{first + 2}" in merges
        assert f"B{first}:B{first + 1}" in merges

    def test_a_side_does_not_merge_across_components(self, sheet):
        # Both machines have a "LADO ACOPLE" and they are not one run.
        merges = {str(merge) for merge in sheet.merged_cells.ranges}
        motor_coupling = _row_of(sheet, "2H")
        pump_coupling = _row_of(sheet, "3H")
        assert f"B{motor_coupling}:B{pump_coupling}" not in merges


def _row_of(sheet, point: str) -> int:
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row=row, column=3).value == point:
            return row
    raise AssertionError(f"point {point} not found")

"""Reads one of IPSA's monthly workbooks into `domain.ipsa.Book`.

Nothing is found by coordinates: the same block sits one column further right
in half the books, and the take and state rows swap places. Every block is
found by the words the analyst wrote on it.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from modules.assets.domain.ipsa import (
    Book,
    Column,
    Limit,
    Note,
    ValueRow,
    book_number,
    component_name,
    limit_of,
    line_speed,
    parse_label,
    plain,
    point_of,
    state_of,
    text_date,
)

SECTIONS = {
    "I. ANTECEDENTES": "background",
    "II. ESTADO ACTUAL": "present",
    "III. RECOMENDACIONES": "recommendation",
}
HEADER = {
    "AREA": "area",
    "EQUIPO": "name",
    "ESTADO": "state",
    "FECHA DE INSP": "inspection",
    "INSPECTOR DE CAMPO": "field",
    "INSPECTOR ANALIS": "analyst",
    "EQUIPO UTILIZADO": "instrument",
}
# Where the operating conditions go: their row names, as the one book that
# records them writes them.
OPERATING = {
    "FRECUENCIA": "freq_hz",
    "PRESION DE SUCCION": "suction_psi",
    "PRESION DE DESCARGA": "discharge_psi",
}


def read_book(path: Path) -> Book:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        grid = _grid(workbook["VIBRACIONES"])
        top, left = _find(grid, "COMPONENTE")
        header = _header(grid[:top], left + 5)
        book = Book(
            filename=path.name,
            number=book_number(path.name) or path.stem,
            area=str(header.get("area") or ""),
            header_name=" ".join(str(header.get("name") or "").split()),
            header_state=state_of(header.get("state")),
            inspection_day=_day(header.get("inspection")),
            analyst=str(header.get("analyst") or "").strip(),
            field_inspector=str(header.get("field") or "").strip(),
            instrument=str(header.get("instrument") or "").strip(),
        )
        _columns(book, grid, top, left)
        _values(book, grid, top, left)
        book.notes.extend(_notes(grid, "vibration", book.inspection_day))
        for name in workbook.sheetnames:
            if plain(name).startswith("TERMOGRAFIA"):
                _thermography(book, _grid(workbook[name]))
    finally:
        workbook.close()
    return book


def _grid(sheet) -> list[list]:
    return [list(row) for row in sheet.iter_rows(max_row=400, values_only=True)]


def _cell(grid, row: int, col: int):
    return grid[row][col] if 0 <= row < len(grid) and 0 <= col < len(grid[row]) else None


def _find(grid, text: str) -> tuple[int, int]:
    for r, row in enumerate(grid):
        for c, value in enumerate(row):
            if isinstance(value, str) and plain(value) == text:
                return r, c
    raise ValueError(f"no '{text}' block")


def _header(grid, right: int) -> dict:
    found: dict = {}
    for row in grid:
        for c, value in enumerate(row[:right]):
            if not isinstance(value, str):
                continue
            key = plain(value).rstrip(".")
            field = next(
                (
                    f
                    for prefix, f in HEADER.items()
                    if key == prefix or (len(prefix) > 6 and key.startswith(prefix))
                ),
                None,
            )
            if field and field not in found:
                found[field] = next((v for v in row[c + 1 : c + 5] if v not in (None, "")), None)
    return found


def _day(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return text_date(value)


def _number(value) -> Decimal | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    try:
        return Decimal(str(value).strip().replace(",", ".")).quantize(Decimal("0.0001"))
    except InvalidOperation:
        return None


def _columns(book: Book, grid, top: int, left: int) -> None:
    width = max(len(row) for row in grid[top : top + 3])
    first, second = top + 1, top + 2
    indexes = [
        c
        for c in range(left + 5, width)
        if any(_cell(grid, r, c) not in (None, "") for r in (top, first, second))
    ]
    # Which of the two rows under the dates is the verdict: the one whose
    # words are states. The other is the take's name ("TOMA 09 VACÍO").
    states_in = {r: sum(1 for c in indexes if state_of(_cell(grid, r, c))) for r in (first, second)}
    state_row = first if states_in[first] >= states_in[second] else second
    take_row = second if state_row == first else first
    for c in indexes:
        raw = _cell(grid, top, c)
        day = _day(raw)
        source = "fecha" if isinstance(raw, (datetime, date)) else ("texto" if day else "")
        state_raw = str(_cell(grid, state_row, c) or "").strip()
        book.columns.append(
            Column(
                index=c,
                day=day,
                source=source,
                take=str(_cell(grid, take_row, c) or "").strip(),
                state=state_of(state_raw),
                raw_state=state_raw,
                speed=line_speed(raw),
            )
        )
        if isinstance(raw, str) and day is not None:
            book.warnings.append(f"fecha escrita como texto: {raw!r} → {day:%d/%m/%Y}")


def _values(book: Book, grid, top: int, left: int) -> None:
    component, point, side = "", None, "custom"
    columns = [c.index for c in book.columns]
    for r in range(top + 3, len(grid)):
        head = _cell(grid, r, left)
        label = _cell(grid, r, left + 4)
        key = plain(head) if isinstance(head, str) else ""
        # Before the end-of-block check: Pulper Nº 04 writes its Parada row on
        # the same line as the "Parámetros de alarma…" caption.
        if isinstance(label, str) and plain(label) in ("PARADA", "ALARMA"):
            limits = book.column_limits.setdefault(plain(label).lower(), set())
            limits.update(v for c in columns if (v := _number(_cell(grid, r, c))) is not None)
            continue
        if key.startswith(("V. LIMITES", "PARAMETROS", "VI. TENDENCIAS")):
            break
        if key.startswith("DATOS DE OPERACION"):
            component = "DATOS DE OPERACIÓN"
        elif isinstance(head, str) and head.strip():
            component = component_name(head.strip())
        if component == "DATOS DE OPERACIÓN":
            name = plain(_cell(grid, r, left + 2))
            code = next((code for prefix, code in OPERATING.items() if name.startswith(prefix)), None)
            if code:
                values = {c: v for c in columns if (v := _number(_cell(grid, r, c))) is not None}
                book.operating[code] = values
            continue
        parsed = parse_label(label)
        if parsed is None:
            if isinstance(head, str) and len(head.strip()) > 30:
                break
            continue
        raw_point = _cell(grid, r, left + 2)
        if raw_point not in (None, ""):
            point, side = point_of(raw_point) or (point, side)
        magnitude, axis, labelled = parsed
        if point is None:
            point = labelled
        if labelled != point:
            book.warnings.append(f"punto {point} rotulado {str(label).strip()}: se guarda como punto {point}")
        row = ValueRow(
            component=component,
            point=point,
            side=side,
            magnitude=magnitude,
            axis=axis,
            label=str(label).strip(),
        )
        row.values = {c: v for c in columns if (v := _number(_cell(grid, r, c))) is not None}
        book.rows.append(row)
    _limits(book, grid, top, left)


def _limits(book: Book, grid, top: int, left: int) -> None:
    for r in range(top, len(grid)):
        title = _cell(grid, r, left)
        if not (isinstance(title, str) and plain(title).startswith("LIMITES PERMISIBLES")):
            continue
        numbers = [n for v in (grid[r + 1] if r + 1 < len(grid) else []) if (n := _number(v)) is not None]
        if len(numbers) >= 2 and limit_of(title):
            book.limits.append(Limit(label=" ".join(title.split()), normal=numbers[0], stop=numbers[-1]))


def _notes(grid, sheet: str, fallback: date | None) -> list[Note]:
    anchors = sorted(
        (r, c, section)
        for r, row in enumerate(grid[:60])
        for c, value in enumerate(row)
        if isinstance(value, str) and (section := SECTIONS.get(plain(value).rstrip(".")))
    )
    notes = []
    for index, (r, c, section) in enumerate(anchors):
        end = anchors[index + 1][0] if index + 1 < len(anchors) else _block_end(grid, r)
        for row in range(r + 1, end):
            # The text sits right beside its date. Looking further right read
            # the registro de valores' header ("10-Sep-25", "T2 - TOMA 06")
            # as recommendations on Pulper Nº 04.
            text = next((v for v in grid[row][c + 1 : c + 3] if isinstance(v, str) and v.strip()), None)
            if text is None or plain(text).startswith(("IV.", "I.", "II.", "III.")):
                continue
            day = _day(_cell(grid, row, c)) or fallback
            if day is not None:
                notes.append(Note(sheet=sheet, section=section, day=day, text=text.strip()))
    return notes


def _block_end(grid, start: int) -> int:
    """The row where "IV. REGISTRO DE VALORES" starts, below the notes."""
    for row in range(start + 1, min(start + 12, len(grid))):
        if any(isinstance(v, str) and plain(v).startswith("IV.") for v in grid[row]):
            return row
    return min(start + 8, len(grid))


def _thermography(book: Book, grid) -> None:
    header = _header(grid[:20], 8)
    day = _day(header.get("inspection"))
    if day and (book.thermo_day is None or day > book.thermo_day):
        book.thermo_day = day
    book.thermo_state = book.thermo_state or state_of(header.get("state"))
    book.notes.extend(_notes(grid, "thermography", day))


def read_summary(path: Path) -> tuple[dict[str, dict], dict[str, dict[str, int]]]:
    """`EQUIPOS - CONCLUSIONES.xlsx`: the verdict per train on 28-sep, and the
    vibration Pareto tables the monthly summary is checked against."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        trains: dict[str, dict] = {}
        for row in workbook.worksheets[0].iter_rows(min_row=2, values_only=True):
            number, name = row[0], row[1]
            if number in (None, "") or not name:
                continue
            key = str(number).strip().upper().replace(" ", "")
            key = key if key.startswith("FDR") else str(int(float(key)))
            trains[key] = {
                "name": " ".join(str(name).split()),
                "thermography": state_of(row[2]),
                "vibration": state_of(row[3]),
                "conclusion": " ".join(str(row[4] or "").split()),
                "recommendation": " ".join(str(row[5] or "").split()),
            }
        pareto: dict[str, dict[str, int]] = {}
        title = None
        for row in workbook["PARETO VIBRACIONAL"].iter_rows(values_only=True):
            first, second = (*row, None, None)[:2]
            if isinstance(second, str):
                title = " ".join(str(first).split())
                pareto[title] = {}
            elif title and state_of(first) and isinstance(second, (int, float)):
                pareto[title][state_of(first)] = int(second)
        return trains, pareto
    finally:
        workbook.close()

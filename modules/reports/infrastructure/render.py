"""One context, three renderings: the screen preview (HTML), the PDF the
customer receives (the same HTML through WeasyPrint), and the workbook."""

from __future__ import annotations

import io
import re

from django.http import HttpResponse
from django.template.loader import render_to_string
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_HEAD = PatternFill("solid", fgColor="1E293B")


def respond(kind: str, context: dict, fmt: str) -> HttpResponse:
    name = _slug(context["filename"])
    if fmt == "xlsx":
        response = HttpResponse(WORKBOOKS[kind](context), content_type=XLSX)
        response["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
        return response
    html = render_to_string(f"reports/{kind}.html", context)
    if fmt == "pdf":
        from weasyprint import HTML

        response = HttpResponse(HTML(string=html).write_pdf(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{name}.pdf"'
        return response
    return HttpResponse(html, content_type="text/html; charset=utf-8")


def _slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "-", text).strip("-").lower() or "informe"


def _sheet(book: Workbook, title: str, header: list[str], rows: list[list], first: bool = False):
    sheet = book.active if first else book.create_sheet()
    sheet.title = title[:31]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        sheet.append([_number(value) for value in row])
    for column in sheet.columns:
        width = max((len(str(cell.value or "")) for cell in column), default=8)
        sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 8), 60)
    return sheet


def _number(value):
    """A thickness stays a number in Excel, so the customer can sort and chart it."""
    if isinstance(value, str) and re.fullmatch(r"-?\d+(\.\d+)?", value):
        return float(value)
    return value


def _save(book: Workbook) -> bytes:
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _mpd(context: dict) -> bytes:
    book = Workbook()
    header = context["header"]
    _sheet(book, "Informe", ["Campo", "Valor"], [
        ["Cliente", header["client"]], ["Locación", header["plant"]], ["Área", header["area"]],
        ["Conjunto", header["group"]], ["Equipos", header["equipments"]],
        ["Estado", header["state"]["name"] if header["state"] else "Sin evaluar"],
        ["Servicio", header["technique"]], ["Orden", header["order"]], ["OT cliente", header["client_order"]],
        ["Fecha", header["dates"]], ["Inspección", header["people"]], ["Analista", header["analyst"]],
        ["Equipo utilizado", header["instruments"]],
    ], first=True)
    _sheet(book, "I-III", ["Sección", "Fecha", "Tipo", "Texto", "Autor", "Estado"], [
        *[["I. Antecedentes", r["date"], r["type"], r["text"], r["author"], ""]
          for r in context["background"]],
        *[["II. Estado actual", r["date"], r["type"], r["text"], r["author"], ""]
          for r in context["present"]],
        *[["III. Recomendaciones", r["date"], r["type"], r["text"], r["author"], r["status"]]
          for r in context["recommendations"]],
    ])
    for block in context["blocks"]:
        _sheet(book, f"IV {block['title']}", ["Componente", "Punto", *context["dates"]],
               [[row["component"], row["label"], *row["values"]] for row in block["rows"]])
    _sheet(book, "V. Límites", ["Magnitud", "Criterio", "Equipos", "Estado", "Desde", "Hasta"], [
        [block["title"], criterion["standard"], ", ".join(criterion["machines"]), band["status"], band["min"],
         band["max"]]
        for block in context["blocks"] for criterion in block["limits"] for band in criterion["bands"]
    ])
    return _save(book)


def _end(context: dict) -> bytes:
    book = Workbook()
    _sheet(book, "Elementos", ["Grupo", "Elemento", *context["labels"], "Titular", "Estado", "Observación"], [
        [row["group"], row["name"], *row["values"], row["headline"], row["state"], row["observation"]]
        for row in context["elements"]
    ], first=True)
    _sheet(book, "Resumen", ["Estado", "Cantidad"],
           [[row["state"], row["count"]] for row in context["summary"]] + [["Total", context["total"]]])
    if context["indications"]:
        _sheet(book, "Incidencias", ["Elemento", "Tipo", "Longitud", "Profundidad", "Posición"],
               [[r["element"], r["kind"], r["length"], r["depth"], r["position"]]
                for r in context["indications"]])
    _sheet(book, "Conclusiones", ["Fecha", "Tipo", "Texto", "Autor"],
           [[r["date"], r["type"], r["text"], r["author"]] for r in context["conclusions"]])
    return _save(book)


def _monthly(context: dict) -> bytes:
    book = Workbook()
    _sheet(book, "Conjuntos", ["Conjunto", *context["techniques"], "Conclusión", "Recomendación"], [
        [row["group"], *[cell["state"] for cell in row["cells"]], row["conclusion"], row["recommendation"]]
        for row in context["rows"]
    ], first=True)
    states = [count["state"] for count in context["pareto"][0]["counts"]] if context["pareto"] else []
    _sheet(book, "Pareto", ["Servicio", *states], [
        [row["technique"], *[count["count"] for count in row["counts"]]] for row in context["pareto"]
    ])
    return _save(book)


def _corrective(context: dict) -> bytes:
    book = Workbook()
    _sheet(book, "Correctivos", ["Conjunto", "Equipo", "Trabajo", "Inicio", "Fin", "Duración", "Responsables",
                                 "Descripción"], [
        [r["group"], r["equipment"], r["work"], r["start"], r["end"], r["duration"], r["responsibles"],
         r["description"]]
        for r in context["rows"]
    ], first=True)
    return _save(book)


WORKBOOKS = {"mpd": _mpd, "end": _end, "monthly": _monthly, "corrective": _corrective}
